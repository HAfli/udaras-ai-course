# -*- coding: utf-8 -*-
"""Build the participant-facing PDF downloads in public/ from the FINAL course materials.

PDF only. Nothing editable, facilitator-only or internal is copied to public/.

  Session-1-Participant-Package.pdf  Workshop 1 (GA, EN) + Workshop 2 (GA, EN) from "Session 1/*.docx" (final, 17 Sept)
  Session-1-Slides.pdf               "Session 1/Udaras_AI_Course_Session_1_MTU_v6.3_DRAFT.pdf" (delivered deck, 55 slides)
  Session-2-Participant-Package.pdf  Quick Guide + Workbook + Fictional Proposal (participant variant, no Organiser's Reply)
  Session-2-Slides.pdf               rendered from "Session 2/Presentation/…pptx" (the delivered, reviewer-edited deck; slides only),
                                     with one Irish correction applied to a temporary copy: "Gan Ceadú" -> "Gan cheadú"
                                     (lenition after gan; FGB). The delivered .pptx itself is not modified.
  Session-3-Participant-Package.pdf  Quick Guide + Workbook + Fictional Scenario
  Session-3-Slides.pdf               "Session 3/Udaras_AI_Course_Session_3_MTU_v1.0_DRAFT.pdf" (40 slides)
  Session-3-Source-Pack.pdf          the eight fictional sources in one PDF (source 4's personal data removed)

Usage: python scripts/build_public_downloads.py   (needs LibreOffice `soffice`, PyMuPDF and python-pptx)
"""
import os
import shutil
import subprocess
import tempfile
import pymupdf

HERE = os.path.dirname(os.path.abspath(__file__))
SITE = os.path.normpath(os.path.join(HERE, ".."))
ROOT = os.path.normpath(os.path.join(SITE, ".."))
PUBLIC = os.path.join(SITE, "public")
S1, S2, S3 = (os.path.join(ROOT, f"Session {n}") for n in (1, 2, 3))

# Public filenames that are produced here; anything else course-related in public/ is removed.
OUTPUTS = ["Session-1-Participant-Package.pdf", "Session-1-Slides.pdf", "Session-2-Participant-Package.pdf", "Session-2-Slides.pdf",
           "Session-3-Participant-Package.pdf", "Session-3-Slides.pdf", "Session-3-Source-Pack.pdf"]
OBSOLETE = ["session1-slides.pdf", "Session_1_Workshop_1_Can_AI_Help_Me.pdf", "Session_1_Workshop_2_Prompt_Challenge.pdf",
            "Session_1_Workshop_1_Can_AI_Help_Me_EN.docx", "Session_1_Workshop_1_Can_AI_Help_Me_GA.docx",
            "Session_1_Workshop_2_Prompt_Challenge_EN.docx", "Session_1_Workshop_2_Prompt_Challenge_GA.docx",
            "Session_3_Fictional_Scenario_GA_EN.pdf", "Session_3_Participant_Quick_Guide_GA_EN.pdf", "Session_3_Participant_Workbook_GA_EN.pdf",
            "Session_3_Source_Pack_All_in_One.pdf", "Session_3_Source_Pack.zip"]


def to_pdf(src, outdir):
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", outdir, src], capture_output=True, timeout=600, check=True)
    return os.path.join(outdir, os.path.splitext(os.path.basename(src))[0] + ".pdf")


S2_SLIDE_FIXES = [("Gan Ceadú", "Gan cheadú")]


def s2_slides_fixed(src, tmp):
    """Copy the delivered Session 2 deck to tmp and apply S2_SLIDE_FIXES to its text runs."""
    from pptx import Presentation
    prs = Presentation(src)
    n = 0
    for slide in prs.slides:
        for sh in slide.shapes:
            if not sh.has_text_frame:
                continue
            for para in sh.text_frame.paragraphs:
                runs = para.runs
                for a, b in S2_SLIDE_FIXES:
                    if a not in "".join(r.text for r in runs):
                        continue
                    # the reviewer's edit left the phrase split across runs ("Gan " + "Ceadú"); fix run by run
                    head, tail = a.split(" ", 1)
                    for k in range(1, len(runs)):
                        if runs[k - 1].text.endswith(head + " ") and runs[k].text.startswith(tail):
                            runs[k].text = b.split(" ", 1)[1] + runs[k].text[len(tail):]
                            n += 1
    assert n == 2, f"expected 2 'Gan Ceadú' runs in the delivered deck, found {n}"
    out = os.path.join(tmp, os.path.basename(src))
    prs.save(out)
    return out


def merge(parts, out):
    doc = pymupdf.open()
    for p in parts:
        doc.insert_pdf(pymupdf.open(p))
    doc.set_metadata({"title": os.path.splitext(os.path.basename(out))[0].replace("-", " "), "author": "MTU × Údarás na Gaeltachta"})
    doc.save(out, garbage=3, deflate=True)
    return out


def main():
    tmp = tempfile.mkdtemp()
    # Session 1
    s1 = [to_pdf(os.path.join(S1, f), tmp) for f in ("Session_1_Workshop_1_Can_AI_Help_Me_GA.docx", "Session_1_Workshop_1_Can_AI_Help_Me_EN.docx",
                                                    "Session_1_Workshop_2_Prompt_Challenge_GA.docx", "Session_1_Workshop_2_Prompt_Challenge_EN.docx")]
    merge(s1, os.path.join(PUBLIC, "Session-1-Participant-Package.pdf"))
    shutil.copyfile(os.path.join(S1, "Udaras_AI_Course_Session_1_MTU_v6.3_DRAFT.pdf"), os.path.join(PUBLIC, "Session-1-Slides.pdf"))
    # Session 2
    pk2 = os.path.join(S2, "Session_2_Participant_Package")
    merge([os.path.join(pk2, f) for f in ("Session_2_Participant_Quick_Guide_GA_EN.pdf", "Session_2_Participant_Workbook_GA_EN.pdf",
                                          "Session_2_Fictional_Proposal_GA_EN.pdf")], os.path.join(PUBLIC, "Session-2-Participant-Package.pdf"))
    s2deck = s2_slides_fixed(os.path.join(S2, "Presentation", "Udaras_AI_Course_Session_2_MTU_v1.0_DRAFT.pptx"), tmp)
    shutil.copyfile(to_pdf(s2deck, tmp), os.path.join(PUBLIC, "Session-2-Slides.pdf"))
    # Session 3
    pk3 = os.path.join(S3, "Session_3_Participant_Package")
    merge([os.path.join(pk3, f) for f in ("Session_3_Participant_Quick_Guide_GA_EN.pdf", "Session_3_Participant_Workbook_GA_EN.pdf",
                                          "Session_3_Fictional_Scenario_GA_EN.pdf")], os.path.join(PUBLIC, "Session-3-Participant-Package.pdf"))
    shutil.copyfile(os.path.join(S3, "Udaras_AI_Course_Session_3_MTU_v1.0_DRAFT.pdf"), os.path.join(PUBLIC, "Session-3-Slides.pdf"))
    shutil.copyfile(os.path.join(S3, "Session_3_Source_Pack", "Session_3_Source_Pack_All_in_One.pdf"), os.path.join(PUBLIC, "Session-3-Source-Pack.pdf"))
    for f in OBSOLETE:
        p = os.path.join(PUBLIC, f)
        if os.path.exists(p):
            os.remove(p)
    shutil.rmtree(tmp, ignore_errors=True)
    for f in OUTPUTS:
        print(f, len(pymupdf.open(os.path.join(PUBLIC, f))), "pages")
    print("public/ now contains:", sorted(x for x in os.listdir(PUBLIC) if not x.startswith(".")))


if __name__ == "__main__":
    main()
